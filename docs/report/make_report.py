# Builds MCSC6002_Final_Report.docx and .pdf from report.md (Word 365 default styling, same as Assignment 2).
# "@table <file> <heading>" pulls the first Markdown table under that heading of another doc into the report.
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
NAME = "MCSC6002_Final_Report"


def set_font(style, name):
    style.font.name = name
    # Theme font attributes override the font name, so drop them.
    rfonts = style.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def add_text(par, text, size=None):
    # `code` spans become Consolas runs on a light grey background.
    for i, part in enumerate(text.split("`")):
        if part:
            run_ = par.add_run(part)
            if size:
                run_.font.size = size
            if i % 2:
                run_.font.name = "Consolas"
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EDEDED")
                run_._r.get_or_add_rPr().append(shd)


def markdown_table(path, heading):
    lines = path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("#") and heading in l)
    rows = []
    for line in lines[start + 1:]:
        if line.startswith("|"):
            if not re.fullmatch(r"\|[-| ]+\|", line):
                rows.append([c.strip() for c in line.strip("|").split("|")])
        elif rows:
            break
    return rows


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            add_text(par, value, Pt(9))
            for run_ in par.runs:
                if run_.font.name is None:
                    run_.font.name = "Aptos"  # the table style's font is the old Calibri theme font
                run_.bold = r == 0
    # Each column gets room for its longest word (no mid-word breaks at 9 pt, ~0.075" per character),
    # and the rest of the 9" width is shared in proportion to the average text length.
    columns = [[row[c] for row in rows] for c in range(len(rows[0]))]
    minimum = [0.075 * max(len(w) for cell in col for w in cell.split() or [""]) + 0.16 for col in columns]
    average = [sum(map(len, col)) / len(col) for col in columns]
    spare = 9 - sum(minimum)
    table.autofit = False
    for c, col in enumerate(columns):
        width = Inches(minimum[c] + spare * average[c] / sum(average))
        for cell in table.columns[c].cells:
            cell.width = width
    header = table.rows[0]._tr.get_or_add_trPr()  # repeat the header row on each page
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)


def build():
    doc = Document()
    # Word 365 defaults; python-docx's template still carries Word 2010 styles.
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme"):  # the old theme font is Calibri
        defaults.rFonts.attrib.pop(qn(attr), None)
    defaults.rFonts.set(qn("w:ascii"), "Aptos")
    defaults.rFonts.set(qn("w:hAnsi"), "Aptos")
    style = doc.styles["Normal"]
    set_font(style, "Aptos")
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15
    for name, font, size, before in [("Heading 1", "Aptos Display", 20, 18),
                                     ("Heading 2", "Aptos Display", 16, 8),
                                     ("Heading 3", "Aptos", 14, 8)]:
        heading = doc.styles[name]
        set_font(heading, font)
        heading.font.size = Pt(size)
        heading.font.bold = False
        heading.font.color.rgb = RGBColor(0x0F, 0x47, 0x61)
        heading.paragraph_format.space_before = Pt(before)
        heading.paragraph_format.space_after = Pt(4)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)

    in_references = False
    for block in (HERE / "report.md").read_text(encoding="utf-8").split("\n\n"):
        block = block.strip()
        if block.startswith("## Appendix A"):
            # Appendices hold wide tables: landscape section.
            new = doc.add_section()
            new.orientation = WD_ORIENT.LANDSCAPE
            new.page_width, new.page_height = Inches(11), Inches(8.5)
        if block.startswith("#"):
            level = len(block) - len(block.lstrip("#"))
            text = block.lstrip("# ")
            doc.add_heading(text, level=min(level, 3))
            in_references = text == "References"
        elif block.startswith("@table"):
            _, path, heading = block.split(" ", 2)
            add_table(doc, markdown_table((HERE / path).resolve(), heading))
        elif block.startswith("!["):
            caption, path = re.fullmatch(r"!\[(.*)\]\((.*)\)", block).groups()
            doc.add_picture(str((HERE / path).resolve()), width=Inches(5.6))  # two figures fit on one page
            doc.add_paragraph(caption).runs[0].italic = True
        elif block.startswith("- "):
            for item in block.splitlines():
                add_text(doc.add_paragraph(style="List Bullet"), item[2:])
        else:
            par = doc.add_paragraph()
            if in_references:  # APA hanging indent
                par.paragraph_format.left_indent = Inches(0.5)
                par.paragraph_format.first_line_indent = Inches(-0.5)
            for i, line in enumerate(block.splitlines()):
                if i:
                    par.add_run().add_break()
                add_text(par, line)
    doc.paragraphs[0].paragraph_format.space_before = Pt(0)  # title starts right at the 1" top margin
    # Replace python-docx's template defaults (author "python-docx", 2013 dates, comment).
    props = doc.core_properties
    props.author = props.last_modified_by = "Arman Akhmetov"
    props.title = "Secure Microservices Inventory System: Final Project Report"
    props.comments = props.subject = props.keywords = ""
    props.created = props.modified = datetime.now(timezone.utc)
    doc.save(HERE / f"{NAME}.docx")


def to_pdf():
    # Word does the conversion so the PDF matches what the .docx looks like.
    docx_path, pdf_path = HERE / f"{NAME}.docx", HERE / f"{NAME}.pdf"
    subprocess.run(["powershell", "-NoProfile", "-Command",
                    f"$w = New-Object -ComObject Word.Application; $d = $w.Documents.Open('{docx_path}'); "
                    f"$d.SaveAs([ref]'{pdf_path}', [ref]17); $d.Close(); $w.Quit()"], check=True)


if __name__ == "__main__":
    build()
    to_pdf()
